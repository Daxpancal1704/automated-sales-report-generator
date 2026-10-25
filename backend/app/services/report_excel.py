import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from typing import Optional, Dict, Any, List
from app.services.analyzer import compute_kpi_stats, compute_trends, compute_product_performance
from app.utils.formatting import get_currency_symbol
from app.utils.security import sanitize_for_spreadsheet

# Modern styling constants
HEADER_FILL = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="0F172A")
SUBTITLE_FONT = Font(name="Calibri", size=11, italic=True, color="64748B")
SECTION_FONT = Font(name="Calibri", size=13, bold=True, color="059669")
BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")
REGULAR_FONT = Font(name="Calibri", size=11, color="1E293B")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0")
)
DOUBLE_BOTTOM_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="059669"),
    bottom=Side(style="double", color="059669")
)

def auto_fit_columns(ws, min_width=12):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if val:
                max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, min_width)

def generate_excel_report(
    df: pd.DataFrame,
    currency: str = "INR",
    validation_summary: Optional[Dict[str, Any]] = None,
    dataset_name: str = "Sales Dataset"
) -> bytes:
    """
    Generate a beautifully styled multi-sheet Excel report with Executive Summary,
    Monthly Performance, Product Performance, Cleaned Sales Data, and Data Quality Audit.
    """
    wb = Workbook()
    currency_sym = get_currency_symbol(currency)
    num_fmt_curr = f'"{currency_sym} "#,##0.00'
    num_fmt_int = '#,##0'
    num_fmt_pct = '0.0%'

    kpis = compute_kpi_stats(df, currency=currency)
    trends = compute_trends(df, grouping="monthly")
    products = compute_product_performance(df)

    # -------------------------------------------------------------
    # Sheet 1: Executive Summary
    # -------------------------------------------------------------
    ws_exec = wb.active
    ws_exec.title = "Executive Summary"
    ws_exec.views.sheetView[0].showGridLines = True

    ws_exec["A1"] = "AUTOMATED SALES REPORT"
    ws_exec["A1"].font = TITLE_FONT
    ws_exec["A2"] = f"Executive Performance Overview | Generated for: {dataset_name}"
    ws_exec["A2"].font = SUBTITLE_FONT

    ws_exec["A4"] = "REPORT METADATA"
    ws_exec["A4"].font = SECTION_FONT
    metadata_rows = [
        ("Reporting Period", f"{kpis.date_range_start or 'N/A'} to {kpis.date_range_end or 'N/A'}"),
        ("Currency Code", f"{currency} ({currency_sym})"),
        ("Unique Orders", kpis.total_orders),
        ("Total Sales Line Items", kpis.total_transactions),
        ("Total Active Products", kpis.total_products_count),
    ]
    for r_idx, (k, v) in enumerate(metadata_rows, start=5):
        ws_exec.cell(row=r_idx, column=1, value=k).font = BOLD_FONT
        ws_exec.cell(row=r_idx, column=2, value=v).font = REGULAR_FONT

    ws_exec["A12"] = "CORE FINANCIAL & OPERATIONAL KPIS"
    ws_exec["A12"].font = SECTION_FONT

    kpi_headers = ["Metric", "Value", "Notes"]
    for c_idx, h in enumerate(kpi_headers, start=1):
        cell = ws_exec.cell(row=13, column=c_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    kpi_items = [
        ("Total Gross Revenue", kpis.total_gross_revenue, num_fmt_curr, "Quantity × Unit Price before discounts"),
        ("Total Discounts Given", kpis.total_discount_amount, num_fmt_curr, "Aggregated customer discounts"),
        ("Total Net Sales Revenue", kpis.total_net_revenue, num_fmt_curr, "Gross Sales minus Discounts"),
        (kpis.order_count_label, kpis.total_orders, num_fmt_int, "Unique order identifiers when available"),
        ("Total Units Sold", kpis.total_units_sold, num_fmt_int, "Sum of all product volumes sold"),
        ("Average Order Value (AOV)", kpis.average_order_value, num_fmt_curr, "Net Revenue divided by Order Count"),
    ]
    if kpis.has_cost_data and kpis.total_cost is not None:
        kpi_items.extend([
            ("Total Cost of Goods Sold", kpis.total_cost, num_fmt_curr, "Quantity × Unit Cost"),
            ("Gross Profit", kpis.total_profit, num_fmt_curr, "Net Revenue minus Total Cost"),
            ("Profit Margin", (kpis.profit_margin_pct or 0.0) / 100.0, num_fmt_pct, "Gross Profit / Net Revenue"),
        ])

    for r_idx, (m_label, m_val, m_fmt, m_note) in enumerate(kpi_items, start=14):
        c1 = ws_exec.cell(row=r_idx, column=1, value=m_label)
        c2 = ws_exec.cell(row=r_idx, column=2, value=m_val)
        c3 = ws_exec.cell(row=r_idx, column=3, value=m_note)
        c1.font = BOLD_FONT
        c2.font = BOLD_FONT
        c3.font = REGULAR_FONT
        c2.number_format = m_fmt
        c1.border = THIN_BORDER
        c2.border = THIN_BORDER
        c3.border = THIN_BORDER
        if r_idx % 2 == 1:
            c1.fill = ZEBRA_FILL
            c2.fill = ZEBRA_FILL
            c3.fill = ZEBRA_FILL

    auto_fit_columns(ws_exec)

    # -------------------------------------------------------------
    # Sheet 2: Monthly Performance + Native Excel Chart
    # -------------------------------------------------------------
    ws_month = wb.create_sheet(title="Monthly Performance")
    ws_month.views.sheetView[0].showGridLines = True
    ws_month.freeze_panes = "A2"

    month_headers = ["Month / Period", "Label", "Net Revenue", "Gross Revenue", "Discounts", "Units Sold", "Transactions", "Orders", "AOV"]
    if kpis.has_cost_data:
        month_headers.append("Gross Profit")

    for c_idx, h in enumerate(month_headers, start=1):
        cell = ws_month.cell(row=1, column=c_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    start_row = 2
    for r_idx, trend in enumerate(trends, start=start_row):
        row_cells = [
            ws_month.cell(row=r_idx, column=1, value=trend.period),
            ws_month.cell(row=r_idx, column=2, value=trend.label),
            ws_month.cell(row=r_idx, column=3, value=trend.net_revenue),
            ws_month.cell(row=r_idx, column=4, value=trend.gross_revenue),
            ws_month.cell(row=r_idx, column=5, value=round(trend.gross_revenue - trend.net_revenue, 2)),
            ws_month.cell(row=r_idx, column=6, value=trend.units_sold),
            ws_month.cell(row=r_idx, column=7, value=trend.transaction_count),
            ws_month.cell(row=r_idx, column=8, value=trend.order_count),
            ws_month.cell(row=r_idx, column=9, value=trend.average_order_value),
        ]
        if kpis.has_cost_data:
            row_cells.append(ws_month.cell(row=r_idx, column=10, value=trend.profit or 0.0))

        for c_cell in row_cells:
            c_cell.font = REGULAR_FONT
            c_cell.border = THIN_BORDER
            if r_idx % 2 == 1:
                c_cell.fill = ZEBRA_FILL

        ws_month.cell(row=r_idx, column=3).number_format = num_fmt_curr
        ws_month.cell(row=r_idx, column=4).number_format = num_fmt_curr
        ws_month.cell(row=r_idx, column=5).number_format = num_fmt_curr
        ws_month.cell(row=r_idx, column=6).number_format = num_fmt_int
        ws_month.cell(row=r_idx, column=7).number_format = num_fmt_int
        ws_month.cell(row=r_idx, column=8).number_format = num_fmt_int
        ws_month.cell(row=r_idx, column=9).number_format = num_fmt_curr
        if kpis.has_cost_data:
            ws_month.cell(row=r_idx, column=10).number_format = num_fmt_curr

    # Totals Row
    last_data_row = start_row + len(trends) - 1
    total_row = last_data_row + 1
    t1 = ws_month.cell(row=total_row, column=1, value="TOTAL")
    t2 = ws_month.cell(row=total_row, column=2, value="")
    t3 = ws_month.cell(row=total_row, column=3, value=f"=SUM(C2:C{last_data_row})")
    t4 = ws_month.cell(row=total_row, column=4, value=f"=SUM(D2:D{last_data_row})")
    t5 = ws_month.cell(row=total_row, column=5, value=f"=SUM(E2:E{last_data_row})")
    t6 = ws_month.cell(row=total_row, column=6, value=f"=SUM(F2:F{last_data_row})")
    t7 = ws_month.cell(row=total_row, column=7, value=f"=SUM(G2:G{last_data_row})")
    t8 = ws_month.cell(row=total_row, column=8, value=f"=SUM(H2:H{last_data_row})")
    t9 = ws_month.cell(row=total_row, column=9, value=f"=C{total_row}/H{total_row}")
    
    for c_cell in [t1, t2, t3, t4, t5, t6, t7, t8, t9]:
        c_cell.font = BOLD_FONT
        c_cell.fill = TOTAL_FILL
        c_cell.border = DOUBLE_BOTTOM_BORDER

    t3.number_format = num_fmt_curr
    t4.number_format = num_fmt_curr
    t5.number_format = num_fmt_curr
    t6.number_format = num_fmt_int
    t7.number_format = num_fmt_int
    t8.number_format = num_fmt_int
    t9.number_format = num_fmt_curr

    # Native Excel Chart
    if len(trends) > 1:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = f"Monthly Net Revenue Trend ({currency_sym})"
        chart.y_axis.title = f"Net Sales ({currency})"
        chart.x_axis.title = "Period"
        data_ref = Reference(ws_month, min_col=3, min_row=1, max_row=last_data_row)
        cats_ref = Reference(ws_month, min_col=2, min_row=2, max_row=last_data_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.height = 12
        chart.width = 18
        ws_month.add_chart(chart, f"L2")

    auto_fit_columns(ws_month)

    # -------------------------------------------------------------
    # Sheet 3: Product Performance
    # -------------------------------------------------------------
    ws_prod = wb.create_sheet(title="Product Performance")
    ws_prod.views.sheetView[0].showGridLines = True
    ws_prod.freeze_panes = "A2"

    prod_headers = ["Rev Rank", "Qty Rank", "Product Name", "Category", "Units Sold", "Net Revenue", "Gross Revenue", "Avg Price", "Revenue Share %"]
    if kpis.has_cost_data:
        prod_headers.extend(["Gross Profit", "Profit Margin %"])

    for c_idx, h in enumerate(prod_headers, start=1):
        cell = ws_prod.cell(row=1, column=c_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for r_idx, prod in enumerate(products, start=2):
        row_cells = [
            ws_prod.cell(row=r_idx, column=1, value=prod.revenue_rank),
            ws_prod.cell(row=r_idx, column=2, value=prod.quantity_rank),
            ws_prod.cell(row=r_idx, column=3, value=sanitize_for_spreadsheet(prod.product_name)),
            ws_prod.cell(row=r_idx, column=4, value=sanitize_for_spreadsheet(prod.category or "General")),
            ws_prod.cell(row=r_idx, column=5, value=prod.units_sold),
            ws_prod.cell(row=r_idx, column=6, value=prod.net_revenue),
            ws_prod.cell(row=r_idx, column=7, value=prod.gross_revenue),
            ws_prod.cell(row=r_idx, column=8, value=prod.avg_unit_price),
            ws_prod.cell(row=r_idx, column=9, value=(prod.revenue_share_pct / 100.0)),
        ]
        if kpis.has_cost_data:
            row_cells.append(ws_prod.cell(row=r_idx, column=10, value=prod.total_profit or 0.0))
            row_cells.append(ws_prod.cell(row=r_idx, column=11, value=((prod.profit_margin_pct or 0.0) / 100.0)))

        for c_cell in row_cells:
            c_cell.font = REGULAR_FONT
            c_cell.border = THIN_BORDER
            if r_idx % 2 == 1:
                c_cell.fill = ZEBRA_FILL

        ws_prod.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
        ws_prod.cell(row=r_idx, column=2).alignment = Alignment(horizontal="center")
        ws_prod.cell(row=r_idx, column=5).number_format = num_fmt_int
        ws_prod.cell(row=r_idx, column=6).number_format = num_fmt_curr
        ws_prod.cell(row=r_idx, column=7).number_format = num_fmt_curr
        ws_prod.cell(row=r_idx, column=8).number_format = num_fmt_curr
        ws_prod.cell(row=r_idx, column=9).number_format = num_fmt_pct
        if kpis.has_cost_data:
            ws_prod.cell(row=r_idx, column=10).number_format = num_fmt_curr
            ws_prod.cell(row=r_idx, column=11).number_format = num_fmt_pct

    auto_fit_columns(ws_prod)

    # -------------------------------------------------------------
    # Sheet 4: Cleaned Sales Data
    # -------------------------------------------------------------
    ws_data = wb.create_sheet(title="Cleaned Sales Data")
    ws_data.views.sheetView[0].showGridLines = True
    ws_data.freeze_panes = "A2"

    data_cols = ["order_id", "date", "product", "category", "quantity", "unit_price", "discount_amount", "gross_sales", "net_sales"]
    if kpis.has_cost_data:
        data_cols.extend(["unit_cost", "total_cost", "profit"])

    col_title_map = {
        "order_id": "Order ID", "date": "Date", "product": "Product", "category": "Category",
        "quantity": "Quantity", "unit_price": "Unit Price", "discount_amount": "Discount Amount",
        "gross_sales": "Gross Sales", "net_sales": "Net Sales", "unit_cost": "Unit Cost",
        "total_cost": "Total Cost", "profit": "Gross Profit"
    }

    for c_idx, col_name in enumerate(data_cols, start=1):
        cell = ws_data.cell(row=1, column=c_idx, value=col_title_map.get(col_name, col_name))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for c_idx, col_name in enumerate(data_cols, start=1):
            val = row.get(col_name)
            cell = ws_data.cell(row=r_idx, column=c_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            if r_idx % 2 == 1:
                cell.fill = ZEBRA_FILL

            if col_name in ["product", "category", "order_id"]:
                cell.value = sanitize_for_spreadsheet(str(val) if pd.notna(val) else "")
            elif col_name == "date":
                cell.value = str(val)
                cell.alignment = Alignment(horizontal="center")
            elif col_name == "quantity":
                cell.value = float(val) if pd.notna(val) else 0
                cell.number_format = num_fmt_int
            else:
                cell.value = float(val) if pd.notna(val) else 0.0
                cell.number_format = num_fmt_curr

    # Enable Excel Auto-Filter on Cleaned Data
    ws_data.auto_filter.ref = f"A1:{get_column_letter(len(data_cols))}{len(df) + 1}"
    auto_fit_columns(ws_data)

    # -------------------------------------------------------------
    # Sheet 5: Data Quality Audit
    # -------------------------------------------------------------
    ws_audit = wb.create_sheet(title="Data Quality Audit")
    ws_audit.views.sheetView[0].showGridLines = True

    ws_audit["A1"] = "DATA INTEGRITY & VALIDATION AUDIT"
    ws_audit["A1"].font = TITLE_FONT

    if validation_summary:
        ws_audit["A3"] = "VALIDATION SUMMARY METRICS"
        ws_audit["A3"].font = SECTION_FONT

        val_metrics = [
            ("Total Raw Records Ingested", validation_summary.get("total_rows", 0)),
            ("Accepted Valid Records", validation_summary.get("accepted_rows", 0)),
            ("Rejected Erroneous Records", validation_summary.get("rejected_rows", 0)),
            ("Data Quality Warnings", validation_summary.get("warning_count", 0)),
        ]
        for r_idx, (m_lbl, m_val) in enumerate(val_metrics, start=4):
            ws_audit.cell(row=r_idx, column=1, value=m_lbl).font = BOLD_FONT
            ws_audit.cell(row=r_idx, column=2, value=m_val).font = REGULAR_FONT
            ws_audit.cell(row=r_idx, column=1).border = THIN_BORDER
            ws_audit.cell(row=r_idx, column=2).border = THIN_BORDER

        reasons = validation_summary.get("rejection_reasons", {})
        if reasons:
            ws_audit["A10"] = "REASONS FOR REJECTED RECORDS"
            ws_audit["A10"].font = SECTION_FONT

            ws_audit.cell(row=11, column=1, value="Rejection Rule / Reason").font = HEADER_FONT
            ws_audit.cell(row=11, column=1).fill = HEADER_FILL
            ws_audit.cell(row=11, column=2, value="Rejected Rows Count").font = HEADER_FONT
            ws_audit.cell(row=11, column=2).fill = HEADER_FILL

            for r_idx, (r_text, r_cnt) in enumerate(reasons.items(), start=12):
                ws_audit.cell(row=r_idx, column=1, value=r_text).font = REGULAR_FONT
                ws_audit.cell(row=r_idx, column=2, value=r_cnt).font = BOLD_FONT
                ws_audit.cell(row=r_idx, column=1).border = THIN_BORDER
                ws_audit.cell(row=r_idx, column=2).border = THIN_BORDER

    auto_fit_columns(ws_audit)

    # Return Excel as binary stream
    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    return excel_stream.getvalue()
