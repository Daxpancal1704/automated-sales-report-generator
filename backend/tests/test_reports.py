import io
import pytest
import pandas as pd
from openpyxl import load_workbook
from app.services.report_excel import generate_excel_report
from app.services.report_pdf import generate_pdf_report
from app.services.report_csv import generate_cleaned_csv, generate_validation_issues_csv

@pytest.fixture
def sample_sales_df():
    return pd.DataFrame([
        {
            "order_id": "ORD-101", "date": "2025-01-15", "month": "2025-01", "year": "2025",
            "product": "Gaming Mouse", "category": "Electronics", "quantity": 2.0, "unit_price": 2500.0,
            "discount_amount": 250.0, "gross_sales": 5000.0, "net_sales": 4750.0,
            "unit_cost": 1500.0, "total_cost": 3000.0, "profit": 1750.0
        },
        {
            "order_id": "ORD-102", "date": "2025-02-20", "month": "2025-02", "year": "2025",
            "product": "Office Chair", "category": "Furniture", "quantity": 1.0, "unit_price": 12000.0,
            "discount_amount": 0.0, "gross_sales": 12000.0, "net_sales": 12000.0,
            "unit_cost": 7000.0, "total_cost": 7000.0, "profit": 5000.0
        }
    ])

def test_generate_excel_report(sample_sales_df):
    val_summary = {
        "total_rows": 2, "accepted_rows": 2, "rejected_rows": 0, "warning_count": 0,
        "rejection_reasons": {}
    }
    excel_bytes = generate_excel_report(sample_sales_df, currency="INR", validation_summary=val_summary)
    
    assert len(excel_bytes) > 0
    # Verify openpyxl can load the generated workbook
    wb = load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Executive Summary" in sheet_names
    assert "Monthly Performance" in sheet_names
    assert "Product Performance" in sheet_names
    assert "Cleaned Sales Data" in sheet_names
    assert "Data Quality Audit" in sheet_names

def test_generate_pdf_report(sample_sales_df):
    val_summary = {
        "total_rows": 2, "accepted_rows": 2, "rejected_rows": 0, "warning_count": 0,
        "rejection_reasons": {}
    }
    pdf_bytes = generate_pdf_report(sample_sales_df, currency="INR", validation_summary=val_summary)
    
    assert len(pdf_bytes) > 0
    # PDF starts with %PDF- header magic bytes
    assert pdf_bytes.startswith(b"%PDF-")

def test_generate_csv_reports(sample_sales_df):
    csv_str = generate_cleaned_csv(sample_sales_df)
    assert "ORDER_ID,DATE,PRODUCT" in csv_str
    assert "Gaming Mouse" in csv_str
    assert "Office Chair" in csv_str

    rejected = [{"row_number": 4, "order_id": "ORD-99", "reason": "Negative quantity", "raw": "data"}]
    err_csv = generate_validation_issues_csv(rejected)
    assert "row_number,order_id,reason" in err_csv
    assert "Negative quantity" in err_csv
